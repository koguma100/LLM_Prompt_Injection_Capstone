import io
from unittest import mock

import openpyxl
import pytest
from docx import Document
from werkzeug.datastructures import FileStorage

from webapp import create_app, routes
from webapp.uploads import UploadError, read_uploads


# A one-page PDF showing text, with a correct cross-reference table
def make_pdf(text):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(b"%PDF-1.4\n"), []
    out.seek(0, 2)
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.write(b"".join(b"%010d 00000 n \n" % offset for offset in offsets))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(text):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Q3", "4.2M"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Quarter", "Revenue"])
    workbook.active.append(["Q3", 4.2])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def upload(name, raw):
    return FileStorage(io.BytesIO(raw), filename=name)


def test_reads_each_supported_type():
    assert read_uploads([upload("report.pdf", make_pdf("Revenue rose 12%."))]) == "Revenue rose 12%."
    assert read_uploads([upload("report.docx", make_docx("Revenue rose."))]) == "Revenue rose.\nQ3\t4.2M"
    assert read_uploads([upload("report.xlsx", make_xlsx())]) == "Quarter\tRevenue\nQ3\t4.2"
    assert read_uploads([upload("notes.TXT", "﻿Plain notes.".encode())]) == "Plain notes."


def test_joins_files_and_skips_the_empty_part_sent_without_a_file():
    files = [upload("a.txt", b"First."), upload("", b""), upload("b.md", b"Second.")]
    assert read_uploads(files) == "First.\n\nSecond."
    assert read_uploads([upload("", b"")]) == ""


@pytest.mark.parametrize("name, raw, message", [
    ("image.png", b"\x89PNG", "isn't a supported file type"),
    ("broken.pdf", b"not a pdf", "Could not read broken.pdf"),
    ("empty.pdf", make_pdf(""), "No text found in empty.pdf"),
])
def test_unusable_files_raise_with_a_message(name, raw, message):
    with pytest.raises(UploadError, match=message):
        read_uploads([upload(name, raw)])


def test_too_large_files_are_rejected():
    with mock.patch("webapp.uploads.MAX_FILE_MB", 0), pytest.raises(UploadError, match="larger than 0 MB"):
        read_uploads([upload("big.txt", b"x")])


@pytest.fixture
def client():
    return create_app().test_client()


def test_running_with_a_pdf_uses_its_text_and_puts_it_in_the_text_box(client):
    with mock.patch.object(routes, "run_experiment", return_value=None) as run:
        html = client.post("/", data={"prompt": "Summarize.", "data": "", "data_file": (io.BytesIO(
            make_pdf("Revenue rose 12%.")), "report.pdf")}, content_type="multipart/form-data").data.decode()
    assert run.call_args.args[:2] == ("Summarize.", "Revenue rose 12%.")
    assert ">Revenue rose 12%.</textarea>" in html


def test_an_unreadable_file_shows_an_error_and_does_not_run(client):
    with mock.patch.object(routes, "run_experiment") as run:
        html = client.post("/", data={"prompt": "Summarize.", "data_file": (io.BytesIO(b"junk"), "x.pdf")},
                           content_type="multipart/form-data").data.decode()
    run.assert_not_called()
    assert "Could not read x.pdf" in html


def test_submitting_without_data_says_what_is_missing(client):
    with mock.patch.object(routes, "run_experiment") as run:
        html = client.post("/", data={"prompt": "Summarize.", "data": ""}).data.decode()
    run.assert_not_called()
    assert "Enter a prompt and some source data" in html


def test_saving_with_a_file_saves_its_text(client):
    client.post("/experiments/save", data={"prompt": "Summarize.", "experiment_name": "From PDF", "data_file": (
        io.BytesIO(make_pdf("Revenue rose 12%.")), "report.pdf")}, content_type="multipart/form-data")
    with create_app().app_context():
        from webapp import experiments
        assert experiments.list_experiments()[0]["data"] == "Revenue rose 12%."
