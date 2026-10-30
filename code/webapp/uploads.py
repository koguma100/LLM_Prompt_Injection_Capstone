# Text from files uploaded as source data on the playground (the data_file input). Uploaded files replace the
# text box, so the pipeline sees only the text extracted here.
import io
import os

import openpyxl
from docx import Document
from pypdf import PdfReader

# Largest file accepted, as the upload area says
MAX_FILE_MB = 20


# A file that can't be used as source data; the message is shown on the page
class UploadError(Exception):
    pass


def _plain_text(raw):
    return raw.decode("utf-8-sig", errors="replace")


# Many PDFs are encrypted with an empty user password (only editing is restricted), so try that first
def _pdf_text(raw):
    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("password-protected PDF")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(raw):
    document = Document(io.BytesIO(raw))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


# One line per non-empty row, cells separated by tabs, every sheet in order
def _xlsx_text(raw):
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cells):
                lines.append("\t".join(cells).rstrip())
    workbook.close()
    return "\n".join(lines)


# The types the file input accepts (playground.html)
READERS = {
    ".txt": _plain_text,
    ".md": _plain_text,
    ".csv": _plain_text,
    ".pdf": _pdf_text,
    ".docx": _docx_text,
    ".xlsx": _xlsx_text,
}


# The text of the uploaded files, separated by blank lines; "" if none were uploaded. Raises UploadError for a file
# that is too large, of another type, unreadable or has no text.
def read_uploads(files):
    texts = []
    for file in files:
        if not file.filename:  # the browser sends an empty part when no file is chosen
            continue
        reader = READERS.get(os.path.splitext(file.filename)[1].lower())
        if reader is None:
            raise UploadError(f"{file.filename} isn't a supported file type. Use PDF, Word (.docx), Excel (.xlsx) "
                              "or a text file.")
        raw = file.read(MAX_FILE_MB * 1024 * 1024 + 1)
        if len(raw) > MAX_FILE_MB * 1024 * 1024:
            raise UploadError(f"{file.filename} is larger than {MAX_FILE_MB} MB.")
        # Parsers raise many kinds of errors on damaged files, and all of them mean the same thing here
        try:
            text = reader(raw).strip()
        except Exception as e:
            raise UploadError(f"Could not read {file.filename}. Check that it isn't damaged or "
                              "password-protected.") from e
        if not text:
            raise UploadError(f"No text found in {file.filename}. Scanned documents (images of text) aren't "
                              "supported.")
        texts.append(text)
    return "\n\n".join(texts)
